"""
Call Register Web App
----------------------
A self-hosted single-service web app: upload your Goal Sheet workbook + a
zip of Sotax service-report PDFs, review the extracted rows in the browser,
edit anything you like (including the blank columns), then download the
finished workbook.

This is intentionally "dumb": it copies fields straight off the report with
no lookups, no guessing, no judgment calls. See README.md for exactly which
columns get filled and which are left blank.

Run locally:
    pip install -r requirements.txt
    uvicorn main:app --reload
    open http://localhost:8000

Deploy: see README.md for one-click deploy notes (Render/Railway/Fly.io).
"""

import copy
import io
import re
import time
import uuid
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import pdfplumber
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

app = FastAPI(title="Call Register Extractor")
app.add_middleware(
    CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"]
)

SHEET_NAME = "Call Register"
LAST_COL = 36  # A .. AJ

# Which columns (1-indexed) are populated straight from the file, and how.
# Everything not listed here is left blank in the preview for manual entry.
AUTO_COLUMNS = {
    1: "sr_no",         # A  Sr. No.                - auto-numbered (not from file content)
    3: "location",      # C  Location               - from filename
    5: "service_type",  # E  Service Type            - from filename
    6: "customer",      # F  Customer Details        - from PDF
    7: "contact",       # G  Contact Person          - from PDF
    9: "email",         # I  Contact Email ID        - from PDF
    10: "model",        # J  Instrument Model        - from PDF (primary description only)
    11: "serial",       # K  Serial Number           - from PDF (first row only)
    16: "ticket_no",    # P  Call Ticket Number      - from PDF ("Service Contract No.")
    17: "ticket_date",  # Q  Call Ticket Date        - from PDF (signature timestamp)
    18: "engineer",     # R  Engineer Allocated      - from PDF ("Service Engineer" field)
}

SERVICE_TYPE_TOKENS = {
    "BD": "Breakdown",
    "AMC": "AMC",
    "PM": "PM",
    "TRAINING": "Training",
    "CV": "CV",
    "LU": "LU",
}

# In-memory store for the workbook between the /extract and /generate calls.
# Keyed by a one-time session id; entries expire after 1 hour.
_SESSIONS = {}
_SESSION_TTL_SECONDS = 3600


def _cleanup_sessions():
    now = time.time()
    expired = [k for k, v in _SESSIONS.items() if now - v["created"] > _SESSION_TTL_SECONDS]
    for k in expired:
        del _SESSIONS[k]


# ---------------------------------------------------------------------------
# PDF extraction (literal fields only - see module docstring)
# ---------------------------------------------------------------------------

def _cluster_rows(words, tolerance=2.5):
    rows = []
    for w in sorted(words, key=lambda w: w["top"]):
        placed = False
        for row in rows:
            if abs(row[0]["top"] - w["top"]) <= tolerance:
                row.append(w)
                placed = True
                break
        if not placed:
            rows.append([w])
    return rows


def extract_pdf_fields(pdf_bytes, filename):
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        page = pdf.pages[0]
        text = page.extract_text(layout=True)
        words = page.extract_words()

    data = {"customer": "", "contact": "", "email": "", "model": "",
            "serial": "", "ticket_no": "", "ticket_date": None, "engineer": ""}

    # --- Customer Information (left column) vs Location of System (right column)
    header_top, boundary_x, service_report_top = None, None, None
    for w in words:
        if w["text"] == "Customer" and header_top is None:
            header_top = w["top"]
        if w["text"] == "Location" and boundary_x is None:
            boundary_x = w["x0"]
    for i, w in enumerate(words):
        if w["text"] == "Service" and i + 1 < len(words) and words[i + 1]["text"] == "Report":
            service_report_top = w["top"]
            break

    if header_top is not None and boundary_x is not None and service_report_top is not None:
        block_words = [w for w in words if header_top < w["top"] < service_report_top]
        rows = sorted(_cluster_rows(block_words), key=lambda row: row[0]["top"])
        left_lines, right_lines = [], []
        for row in rows:
            row_sorted = sorted(row, key=lambda w: w["x0"])
            left = " ".join(w["text"] for w in row_sorted if w["x0"] < boundary_x - 1)
            right = " ".join(w["text"] for w in row_sorted if w["x0"] >= boundary_x - 1)
            if left.strip():
                left_lines.append(left.strip())
            if right.strip():
                right_lines.append(right.strip())

        if left_lines:
            data["customer"] = left_lines[0]
        for l in right_lines:
            if re.match(r"^(Mr\.?|Ms\.?|Mrs\.?)\s*\S", l) and "@" not in l:
                data["contact"] = l
                break
        for l in right_lines:
            m = re.search(r"[\w.\-]+@[\w.\-]+\.\w+", l)
            if m:
                data["email"] = m.group(0)
                break

    # --- Service Contract No (ticket number)
    m = re.search(r"Service Contract No\.\s+(\S+)", text)
    if m:
        data["ticket_no"] = m.group(1)

    # --- Service Engineer (strip any leading numeric code, e.g. "00000Rahul Sharma")
    m = re.search(r"Service Engineer\s+(?:\d+)?([A-Za-z][A-Za-z .]*?)\s+Fault Reason", text)
    if m:
        data["engineer"] = m.group(1).strip()

    # --- Signature date/time (earliest of the footer timestamps)
    dt_matches = re.findall(r"Date:\s*(\d{2})\.(\d{2})\.(\d{4})\s+(\d{2}):(\d{2})", text)
    dts = []
    for d, mo, y, h, mi in dt_matches:
        try:
            dts.append(datetime(int(y), int(mo), int(d), int(h), int(mi)))
        except ValueError:
            pass
    if dts:
        data["ticket_date"] = min(dts)

    # --- Serial Number / Instrument Description table (first row only)
    lines = text.split("\n")
    capture = False
    for l in lines:
        if "Serial Number" in l and "Instrument Description" in l:
            capture = True
            continue
        if capture:
            if l.strip() == "" or "VISA1" in l:
                break
            parts = re.split(r"\s{2,}", l.strip())
            if parts:
                data["serial"] = parts[0]
                data["model"] = parts[1] if len(parts) > 1 else ""
            break

    return data


def parse_filename(filename):
    """{Company}_{Location}_{Model}_{Serial}_{Type}_...date.pdf"""
    stem = Path(filename).stem
    tokens = stem.split("_")
    location = tokens[1] if len(tokens) > 1 else ""

    service_type = ""
    stem_upper = stem.upper()
    for token, mapped in SERVICE_TYPE_TOKENS.items():
        if re.search(rf"(?<![A-Z]){token}(?![A-Z])", stem_upper):
            service_type = mapped
            break

    return location, service_type


def find_last_data_row(ws):
    last = 0
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is not None:
            last = r
    return last


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------

@app.post("/api/extract")
async def extract(xlsx: UploadFile = File(...), zipfile_: UploadFile = File(..., alias="zip")):
    _cleanup_sessions()

    xlsx_bytes = await xlsx.read()
    zip_bytes = await zipfile_.read()

    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    except Exception:
        raise HTTPException(400, "Could not read the uploaded workbook. Is it a valid .xlsx file?")

    if SHEET_NAME not in wb.sheetnames:
        raise HTTPException(400, f"No sheet named '{SHEET_NAME}' found in the workbook.")

    ws = wb[SHEET_NAME]
    last_row = find_last_data_row(ws)
    last_sr = ws.cell(row=last_row, column=1).value
    next_sr = (last_sr + 1) if isinstance(last_sr, (int, float)) else last_row

    headers = [ws.cell(row=1, column=c).value or "" for c in range(1, LAST_COL + 1)]

    try:
        zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    except zipfile.BadZipFile:
        raise HTTPException(400, "That doesn't look like a valid zip file.")

    pdf_names = sorted(n for n in zf.namelist() if n.lower().endswith(".pdf"))
    if not pdf_names:
        raise HTTPException(400, "No PDF files found inside the zip.")

    parsed = []
    for name in pdf_names:
        pdf_bytes = zf.read(name)
        try:
            fields = extract_pdf_fields(pdf_bytes, name)
        except Exception as e:
            fields = {"customer": "", "contact": "", "email": "", "model": "",
                      "serial": "", "ticket_no": "", "ticket_date": None, "engineer": ""}
        location, service_type = parse_filename(name)
        fields["location"] = location
        fields["service_type"] = service_type
        fields["_filename"] = Path(name).name
        parsed.append(fields)

    # Chronological order (rows with no detected date sort last, in filename order)
    parsed.sort(key=lambda f: (f["ticket_date"] is None, f["ticket_date"] or datetime.max))

    rows = []
    for i, fields in enumerate(parsed):
        row = {str(c): "" for c in range(1, LAST_COL + 1)}
        row["1"] = next_sr + i
        row["3"] = fields["location"]
        row["5"] = fields["service_type"]
        row["6"] = fields["customer"]
        row["7"] = fields["contact"]
        row["9"] = fields["email"]
        row["10"] = fields["model"]
        row["11"] = fields["serial"]
        row["16"] = fields["ticket_no"]
        row["17"] = fields["ticket_date"].strftime("%Y-%m-%d %H:%M") if fields["ticket_date"] else ""
        row["18"] = fields["engineer"]
        row["_filename"] = fields["_filename"]
        rows.append(row)

    session_id = str(uuid.uuid4())
    _SESSIONS[session_id] = {"xlsx_bytes": xlsx_bytes, "created": time.time()}

    return {
        "session_id": session_id,
        "headers": headers,
        "auto_columns": list(AUTO_COLUMNS.keys()),
        "rows": rows,
        "sheet_name": SHEET_NAME,
    }


class GenerateRequest(BaseModel):
    session_id: str
    rows: list[dict]


@app.post("/api/generate")
async def generate(req: GenerateRequest):
    session = _SESSIONS.get(req.session_id)
    if not session:
        raise HTTPException(400, "This session has expired. Please re-upload your files.")

    wb = openpyxl.load_workbook(io.BytesIO(session["xlsx_bytes"]), data_only=False)
    ws = wb[SHEET_NAME]
    last_row = find_last_data_row(ws)
    start_row = last_row + 1

    DATE_COLS = {17}  # Q - Call Ticket Date

    for i, row_data in enumerate(req.rows):
        r = start_row + i
        for col in range(1, LAST_COL + 1):
            src_cell = ws.cell(row=last_row, column=col)
            ws.cell(row=r, column=col)._style = copy.copy(src_cell._style)

        for col in range(1, LAST_COL + 1):
            raw_value = row_data.get(str(col), "")
            cell = ws.cell(row=r, column=col)

            if raw_value == "" or raw_value is None:
                cell.value = None
                continue

            if col == 1:
                try:
                    cell.value = int(raw_value)
                except (ValueError, TypeError):
                    cell.value = raw_value
            elif col in DATE_COLS:
                parsed_dt = None
                for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M", "%d.%m.%Y %H:%M", "%d/%m/%Y %H:%M"):
                    try:
                        parsed_dt = datetime.strptime(str(raw_value), fmt)
                        break
                    except ValueError:
                        continue
                cell.value = parsed_dt if parsed_dt else raw_value
            elif col == 11:  # Serial Number - keep as text so leading digits aren't lost
                cell.value = str(raw_value)
                cell.number_format = "@"
            else:
                cell.value = raw_value

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    del _SESSIONS[req.session_id]

    tmp_path = Path(f"/tmp/call_register_{uuid.uuid4().hex}.xlsx")
    tmp_path.write_bytes(out.getvalue())

    return FileResponse(
        tmp_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="Goal_Sheet_updated.xlsx",
    )


# ---------------------------------------------------------------------------
# Static frontend
# ---------------------------------------------------------------------------

static_dir = Path(__file__).parent / "static"
app.mount("/static", StaticFiles(directory=static_dir), name="static")


@app.get("/", response_class=HTMLResponse)
async def index():
    return (static_dir / "index.html").read_text()
